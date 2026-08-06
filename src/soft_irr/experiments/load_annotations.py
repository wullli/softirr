
from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable, Iterator, Literal

from tqdm import tqdm

from soft_irr.common import load_dotenv
from soft_irr.experiments.load_taxonomies import ConceptEntry

CategoryField = Literal["disease_label", "skin_concept", "both"]
_DERM1_CATEGORY_FIELDS = ("disease_label", "skin_concept")
PassionCategoryField = Literal["diagnosis", "conditions_PASSION"]

_PASSION_DESCRIPTION_COLUMNS = ("descr1", "descr2", "descr3", "descr4")

_MISSING_CATEGORY_VALUES = {
    "disease_label": {"no definitive diagnosis"},
    "skin_concept": {"no visual concepts"},
}

_REFLACX_CERTAINTY_LABELS: dict[int, list[str]] = {
    1: [
        "Airway wall thickening", "Atelectasis", "Consolidation", "Emphysema",
        "Enlarged cardiac silhouette", "Fibrosis", "Fracture", "Groundglass opacity",
        "Mass", "Nodule", "Pleural effusion", "Pleural thickening", "Pneumothorax",
        "Pulmonary edema", "Wide mediastinum",
    ],
    2: [
        "Abnormal mediastinal contour", "Acute fracture", "Atelectasis", "Consolidation",
        "Enlarged cardiac silhouette", "Enlarged hilum", "Groundglass opacity", "Hiatal hernia",
        "High lung volume / emphysema", "Interstitial lung disease", "Lung nodule or mass",
        "Pleural abnormality", "Pneumothorax", "Pulmonary edema",
    ],
}
_REFLACX_BOOLEAN_LABELS: dict[int, list[str]] = {
    1: ["Quality issue", "Support devices"],
    2: ["Support devices"],
}

_ROOT_DIR = Path(__file__).parents[3]
load_dotenv(str(_ROOT_DIR / ".env"))


def _image_hash(image: Any) -> str:
    import imagehash

    return str(imagehash.phash(image))


def _load_hash_cache(cache_path: Path) -> dict[str, str]:
    if not cache_path.exists():
        return {}
    with open(cache_path, newline="", encoding="utf-8") as f:
        return {row[0]: row[1] for row in csv.reader(f) if len(row) == 2}


def _append_hash_cache(cache_path: Path, key: str, phash: str) -> None:
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    with open(cache_path, "a", newline="", encoding="utf-8") as f:
        csv.writer(f).writerow([key, phash])


def _iter_zip_images(repo_id: str, archive: str, wanted: dict[str, str], tmp_dir: Path):
    import zipfile

    from huggingface_hub import hf_hub_download
    from PIL import Image as PILImage

    tmp_dir.mkdir(parents=True, exist_ok=True)
    local_path = Path(
        hf_hub_download(repo_id=repo_id, repo_type="dataset", filename=archive, local_dir=str(tmp_dir))
    )
    try:
        with zipfile.ZipFile(local_path) as zf:
            names_in_zip = set(zf.namelist())
            for member_name in wanted.keys() & names_in_zip:
                try:
                    with zf.open(member_name) as fh:
                        image = PILImage.open(fh)
                        image.load()
                except Exception as exc:
                    tqdm.write(f"  [WARN] Failed to decode {archive}/{member_name}: {exc}")
                    continue
                yield wanted[member_name], image
    finally:
        local_path.unlink(missing_ok=True)


def _derm1_parent(hierarchical_disease_label: str, disease_label: str) -> str:
    parts = [p.strip() for p in hierarchical_disease_label.split(",") if p.strip()]
    parts = [p for p in parts if p.casefold() != disease_label.casefold()]
    return ", ".join(parts)


def _derm1_wanted_members(filenames: Iterable[str], hash_cache: dict[str, str]) -> dict[str, str]:
    pending = [filename for filename in filenames if filename not in hash_cache]
    wanted: dict[str, str] = {filename: filename for filename in pending}
    for filename in pending:
        wanted.setdefault(filename.split("/", 1)[1], filename)
    return wanted


def _iter_derm1_hashed_rows(
    cache_dir: str | None = None,
    hash_cache_path: str | Path | None = None,
    max_rows: int | None = None,
    archive_tmp_dir: str | Path | None = None,
) -> Iterator[tuple[str, dict[str, Any]]]:
    from datasets import IterableDataset as HFIterableDataset
    from datasets import load_dataset

    ds = load_dataset("redlessone/Derm1M", split="train", streaming=True, cache_dir=cache_dir)
    assert isinstance(ds, HFIterableDataset)
    ds = ds.select_columns(
        ["filename", "caption", "disease_label", "hierarchical_disease_label", "skin_concept"]
    )
    if max_rows is not None:
        ds = ds.take(max_rows)

    hash_cache_path = Path(hash_cache_path) if hash_cache_path else _ROOT_DIR / "data" / "derm1_image_hashes.csv"
    hash_cache = _load_hash_cache(hash_cache_path)
    archive_tmp_dir = Path(archive_tmp_dir) if archive_tmp_dir else _ROOT_DIR / "data" / "_derm1_archive_tmp"

    rows_by_archive: dict[str, dict[str, dict[str, str]]] = defaultdict(dict)
    for _row in tqdm(ds, desc="Scanning Derm1M metadata", unit="row"):
        row: dict[str, Any] = dict(_row)
        filename = (row.get("filename") or "").strip()
        caption = (row.get("caption") or "").strip()
        if not filename or not caption or "/" not in filename:
            continue
        archive = filename.split("/", 1)[0] + ".zip"
        rows_by_archive[archive][filename] = row

    for archive, by_filename in tqdm(rows_by_archive.items(), desc="Derm1M archives", unit="archive"):
        wanted = _derm1_wanted_members(by_filename, hash_cache)
        if wanted:
            for filename, image in _iter_zip_images("redlessone/Derm1M", archive, wanted, archive_tmp_dir):
                try:
                    phash = _image_hash(image)
                except Exception as exc:
                    tqdm.write(f"  [WARN] Failed to hash {filename}: {exc}")
                    continue
                hash_cache[filename] = phash
                _append_hash_cache(hash_cache_path, filename, phash)

        for filename, row in by_filename.items():
            phash = hash_cache.get(filename)
            if phash is None:
                continue
            yield phash, row


def load_derm1(
    cache_dir: str | None = None,
    hash_cache_path: str | Path | None = None,
    max_rows: int | None = None,
    archive_tmp_dir: str | Path | None = None,
) -> dict[str, ConceptEntry]:
    captions_by_hash: dict[str, list[str]] = defaultdict(list)
    parent_by_hash: dict[str, str] = {}
    seen: dict[str, set[str]] = defaultdict(set)

    for phash, row in _iter_derm1_hashed_rows(cache_dir, hash_cache_path, max_rows, archive_tmp_dir):
        caption = row["caption"].strip()
        key = caption.casefold()
        if key not in seen[phash]:
            seen[phash].add(key)
            captions_by_hash[phash].append(caption)

        disease = (row.get("disease_label") or "").strip()
        hierarchy = (row.get("hierarchical_disease_label") or "").strip()
        if disease and phash not in parent_by_hash:
            parent = _derm1_parent(hierarchy, disease) if hierarchy else ""
            if parent:
                parent_by_hash[phash] = parent

    concepts: dict[str, ConceptEntry] = {}
    for phash, captions in captions_by_hash.items():
        if len(captions) < 2:
            continue
        preferred, *synonyms = captions
        parent = parent_by_hash.get(phash, "")
        concepts[phash] = ConceptEntry(
            code=phash,
            preferred_term=preferred,
            synonyms=synonyms,
            parent_codes=[parent] if parent else [],
        )
    return concepts


def _derm1_category(row: dict[str, Any], category_field: CategoryField) -> frozenset[str] | None:
    if category_field == "both":
        combined: set[str] = set()
        for field in _DERM1_CATEGORY_FIELDS:
            part = _derm1_category(row, field)
            if part is not None:
                combined.update(f"{field}:{t}" for t in part)
        return frozenset(combined) or None
    raw = (row.get(category_field) or "").strip()
    if not raw or raw.casefold() in _MISSING_CATEGORY_VALUES[category_field]:
        return None
    if category_field == "disease_label":
        return frozenset({raw.casefold()})
    tags = frozenset(t.strip().casefold() for t in raw.split(",") if t.strip())
    return tags or None


def load_derm1_annotation_rows(
    category_field: CategoryField,
    cache_dir: str | None = None,
    hash_cache_path: str | Path | None = None,
    max_rows: int | None = None,
    archive_tmp_dir: str | Path | None = None,
) -> dict[str, list[tuple[str, frozenset[str], str]]]:
    rows_by_hash: dict[str, list[tuple[str, frozenset[str], str]]] = defaultdict(list)
    for phash, row in _iter_derm1_hashed_rows(cache_dir, hash_cache_path, max_rows, archive_tmp_dir):
        category = _derm1_category(row, category_field)
        if category is None:
            continue
        archive = row["filename"].split("/", 1)[0]
        rows_by_hash[phash].append((row["caption"].strip(), category, archive))
    return {
        phash: rows for phash, rows in rows_by_hash.items()
        if len({archive for _, _, archive in rows}) >= 2
    }


def _reflacx_category(row: dict[str, Any], phase: int) -> frozenset[str]:
    tags: set[str] = set()
    for label in _REFLACX_CERTAINTY_LABELS[phase]:
        raw = (row.get(label) or "").strip()
        try:
            if raw and int(raw) > 0:
                tags.add(label.casefold())
        except ValueError:
            continue
    for label in _REFLACX_BOOLEAN_LABELS[phase]:
        if (row.get(label) or "").strip().casefold() == "true":
            tags.add(label.casefold())
    return frozenset(tags)


def _iter_reflacx_rows(reflacx_dir: Path, phase: int, max_rows: int | None) -> Iterator[dict[str, Any]]:
    metadata_path = reflacx_dir / "main_data" / f"metadata_phase_{phase}.csv"
    with open(metadata_path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for i, row in enumerate(reader):
            if max_rows is not None and i >= max_rows:
                break
            if row.get("eye_tracking_data_discarded", "False").strip().casefold() == "true":
                continue
            yield row


def load_reflacx_annotation_rows(
    reflacx_dir: str | Path | None = None,
    phases: tuple[int, ...] = (1, 2),
    max_rows: int | None = None,
) -> dict[str, list[tuple[str, frozenset[str], str]]]:
    reflacx_dir = Path(reflacx_dir) if reflacx_dir else _ROOT_DIR / "data" / "reflacx"

    rows_by_dicom: dict[str, list[tuple[str, frozenset[str], str]]] = defaultdict(list)
    dicom_ids_by_phase: dict[int, set[str]] = {}
    for phase in phases:
        seen_this_phase: set[str] = set()
        for row in tqdm(
            _iter_reflacx_rows(reflacx_dir, phase, max_rows), desc=f"Scanning REFLACX phase {phase}", unit="row"
        ):
            reading_id = row["id"].strip()
            transcription_path = reflacx_dir / "main_data" / reading_id / "transcription.txt"
            try:
                transcription = transcription_path.read_text(encoding="utf-8").strip()
            except FileNotFoundError:
                continue
            if not transcription:
                continue
            category = _reflacx_category(row, phase)
            dicom_id = row["dicom_id"].strip()
            seen_this_phase.add(dicom_id)
            rows_by_dicom[dicom_id].append((transcription, category, reading_id))
        dicom_ids_by_phase[phase] = seen_this_phase

    for i, phase_a in enumerate(phases):
        for phase_b in phases[i + 1:]:
            overlap = sorted(dicom_ids_by_phase[phase_a] & dicom_ids_by_phase[phase_b])
            if overlap:
                shown = ", ".join(overlap[:10]) + (" ..." if len(overlap) > 10 else "")
                raise ValueError(
                    f"REFLACX phases {phase_a} and {phase_b} share {len(overlap):,} dicom_id(s), "
                    f"but use disjoint label schemas, so their readings cannot be pooled: {shown}"
                )

    return {dicom_id: rows for dicom_id, rows in rows_by_dicom.items() if len(rows) >= 2}


def _load_passion_labels(labels_path: Path) -> dict[str, dict[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(labels_path, read_only=True, data_only=True)
    ws = wb["PASSION_cleaned_final"]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    idx = {name: i for i, name in enumerate(header)}

    labels: dict[str, dict[str, str]] = {}
    for row in rows:
        subject_id = row[idx["subject_id"]]
        if subject_id is None:
            continue
        labels[str(subject_id).strip()] = {
            "diagnosis": str(row[idx["diagnosis"]] or "").strip(),
            "conditions_PASSION": str(row[idx["conditions_PASSION"]] or "").strip(),
        }
    return labels


def load_passion_annotation_rows(
    category_field: PassionCategoryField,
    passion_dir: str | Path | None = None,
    max_rows: int | None = None,
) -> dict[str, list[tuple[str, frozenset[str], str]]]:
    passion_dir = Path(passion_dir) if passion_dir else _ROOT_DIR / "data" / "passion"
    labels = _load_passion_labels(passion_dir / "passion_labels.xlsx")

    rows_by_image: dict[str, list[tuple[str, frozenset[str], str]]] = {}
    with open(passion_dir / "passion_descriptions.csv", newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for i, row in enumerate(reader):
            if max_rows is not None and i >= max_rows:
                break
            img_path = (row.get("img_path") or "").strip()
            if not img_path:
                continue
            subject_id = img_path.rsplit("_", 1)[0]
            label = labels.get(subject_id, {}).get(category_field, "")
            if not label:
                continue
            category = frozenset({label.casefold()})
            raters = [
                (row[col].strip(), category, col)
                for col in _PASSION_DESCRIPTION_COLUMNS
                if (row.get(col) or "").strip()
            ]
            if len(raters) >= 2:
                rows_by_image[img_path] = raters

    return rows_by_image
